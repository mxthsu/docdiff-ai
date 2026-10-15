"""Profile small synthetic conversion inputs without downloading ML models."""

from __future__ import annotations

import json
import tempfile
import time
import tracemalloc
from pathlib import Path

from docdiff_ai.documents import from_path


def generate(root: Path) -> list[Path]:
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    from reportlab.pdfgen import canvas

    lines = [
        "Total price: USD 1250.00",
        "Delivery time: 30 days",
        "Warranty: 12 months",
        "Payment terms: 15 days",
    ]
    pdf = root / "sample.pdf"
    page = canvas.Canvas(str(pdf))
    for index, line in enumerate(lines):
        page.drawString(70, 730 - 28 * index, line)
    page.save()

    docx = root / "sample.docx"
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    document.save(docx)

    pptx = root / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4))
    textbox.text_frame.text = lines[0]
    for line in lines[1:]:
        textbox.text_frame.add_paragraph().text = line
    presentation.save(pptx)
    return [pdf, docx, pptx]


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="docdiff-profile-") as folder:
        for path in generate(Path(folder)):
            tracemalloc.start()
            cpu_start = time.process_time()
            wall_start = time.perf_counter()
            document = from_path(path)
            seconds = time.perf_counter() - wall_start
            cpu_seconds = time.process_time() - cpu_start
            _current, peak = tracemalloc.get_traced_memory()
            tracemalloc.stop()
            result = {
                "format": path.suffix.lower(),
                "input_bytes": path.stat().st_size,
                "output_chars": len(document.text),
                "wall_seconds": round(seconds, 3),
                "cpu_seconds": round(cpu_seconds, 3),
                "tracemalloc_python_peak_mib": round(peak / 1024**2, 2),
                "caveat": "Peak covers Python allocations only, not full native RSS.",
            }
            print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
