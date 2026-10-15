"""Slow, optional end-to-end tests for actual Docling PDF and DOCX conversion.

Run explicitly: python -m unittest discover -s tests/integration -v
Requires: pip install -e '.[documents,integration-tests]'
"""

import tempfile
import unittest
from pathlib import Path

from docdiff_ai.core import compare, extract_evidence
from docdiff_ai.documents import DocumentReadError, from_bytes, from_path

try:
    from docx import Document as WordDocument
    from pptx import Presentation
    from pptx.util import Inches
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
except ImportError:
    WordDocument = None
    canvas = None
    Presentation = None
    Inches = None


@unittest.skipUnless(WordDocument is not None and canvas is not None, "Missing document-test extras")
class DoclingConversionIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="docdiff-integration-")
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)

    def make_pdf(self, filename: str, price: str = "USD 1250.00") -> Path:
        path = self.root / filename
        page = canvas.Canvas(str(path), pagesize=letter)
        lines = [
            f"Total price: {price}",
            "Delivery time: 30 days",
            "Warranty: 12 months",
            "Payment terms: 15 days",
        ]
        for index, line in enumerate(lines):
            page.drawString(72, 700 - (index * 30), line)
        page.save()
        return path

    def make_docx(self, filename: str, price: str = "USD 1250.00") -> Path:
        path = self.root / filename
        doc = WordDocument()
        for line in [
            f"Total price: {price}",
            "Delivery time: 30 days",
            "Warranty: 12 months",
            "Payment terms: 15 days",
        ]:
            doc.add_paragraph(line)
        doc.save(path)
        return path

    def make_pptx(self, filename: str) -> Path:
        path = self.root / filename
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(4)).text_frame
        frame.text = "Total price: USD 1250.00"
        for line in ("Delivery time: 30 days", "Warranty: 12 months", "Payment terms: 15 days"):
            frame.add_paragraph().text = line
        presentation.save(path)
        return path

    def test_pptx_conversion_preserves_values(self):
        doc = from_path(self.make_pptx("proposal.pptx"))
        self.assertEqual(doc.reference_type, "converted_markdown_line")
        self.assertEqual(extract_evidence(doc)["total_price"][0].normalized, "USD:1250.00")

    def test_pdf_native_pipeline_extracts_without_models(self):
        document = from_path(self.make_pdf("proposal.pdf"))
        self.assertEqual(document.reference_type, "converted_markdown_line")
        hits = extract_evidence(document)
        self.assertEqual(hits["total_price"][0].normalized, "USD:1250.00")
        self.assertEqual(hits["delivery_days"][0].normalized, "30")
        self.assertGreater(hits["total_price"][0].line, 0)
        self.assertEqual(hits["total_price"][0].reference_type, "converted_markdown_line")

    def test_docx_conversion_preserves_text_evidence(self):
        document = from_path(self.make_docx("proposal.docx"))
        self.assertEqual(document.reference_type, "converted_markdown_line")
        hits = extract_evidence(document)
        self.assertEqual(hits["total_price"][0].normalized, "USD:1250.00")
        self.assertEqual(hits["warranty_months"][0].normalized, "12")

    def test_two_real_formats_compare(self):
        pdf = from_path(self.make_pdf("alpha.pdf"))
        word = from_path(self.make_docx("beta.docx", price="USD 1400.00"))
        result = compare([pdf, word])
        self.assertEqual(result.fields[0].status, "different")
        self.assertEqual(result.fields[1].status, "match")
        for hit in result.fields[0].evidence["alpha.pdf"]:
            self.assertEqual(hit.reference_type, "converted_markdown_line")

    def test_upload_bytes_use_safe_temporary_path(self):
        filename = self.make_docx("temp.docx")
        document = from_bytes("../unsafe.docx", filename.read_bytes())
        self.assertEqual(document.name, "unsafe.docx")
        self.assertEqual(document.reference_type, "converted_markdown_line")
        self.assertIn("Total price", document.text)

    def test_corrupted_pdf_is_rejected(self):
        path = self.root / "corrupted.pdf"
        path.write_bytes(b"%PDF-1.7\nnot a real PDF")
        with self.assertRaises(DocumentReadError):
            from_path(path)

    def test_empty_docx_is_rejected(self):
        path = self.root / "empty.docx"
        WordDocument().save(path)
        with self.assertRaises(DocumentReadError):
            from_path(path)


if __name__ == "__main__":
    unittest.main()
